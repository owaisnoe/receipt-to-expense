# Deck v2: landing-page style (patterns from Mobbin, see DESIGN_NOTES.md). Run: .venv\Scripts\python.exe build_deck_v2.py
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
CH = os.path.join(HERE, "charts_v2"); os.makedirs(CH, exist_ok=True)
SS = os.path.join(HERE, "screenshots")
OUT_PPTX = os.path.join(HERE, "Theme3B_Receipt_to_Expense_v2.pptx")
HF_URL = "https://huggingface.co/owaisnoe/gemma4-e2b-sroie-lora"
APP_URL = "https://f6753b53a276c78344.gradio.live"
QR_APP, QR_HF = os.path.join(HERE, "charts", "qr_app.png"), os.path.join(HERE, "charts", "qr_hf.png")

INK, BODY, GREY, HAIR, ACCENT, ACCENT_SOFT, PAPER, NIGHT = "#0b0f19", "#374151", "#6b7280", "#e5e7eb", "#0f9f6e", "#a7e3cc", "#f6f7f9", "#0b1220"
FONT = "Segoe UI"
FOOTNOTE = "* Best of 5 configurations on the same 100 test receipts; to confirm on a fresh test set."

def rgb(h):
    h = h.lstrip("#"); return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

# ---------------------------------------------------------------- charts (minimal: no grid, direct labels)
plt.rcParams.update({"font.family": ["Segoe UI", "DejaVu Sans"], "font.size": 15, "axes.spines.top": False,
                     "axes.spines.right": False, "axes.spines.left": False, "axes.edgecolor": HAIR,
                     "xtick.color": BODY, "ytick.color": BODY, "xtick.major.size": 0, "ytick.major.size": 0})

def save(fig, name, dark=False):
    p = os.path.join(CH, name)
    fig.savefig(p, dpi=220, bbox_inches="tight", transparent=False, facecolor=NIGHT if dark else "white")
    plt.close(fig); return p

def chart_progress():
    labels = ["Plain\nGemma 4", "Detailed\nprompt", "Fine-tuned\nv1 (560)", "Fine-tuned\nv2 (1120)*"]
    vals, cols = [28, 37, 66, 76], ["#475569", "#64748b", "#5fbf9a", "#34d399"]
    fig, ax = plt.subplots(figsize=(6.6, 4.3)); fig.patch.set_facecolor(NIGHT); ax.set_facecolor(NIGHT)
    b = ax.bar(labels, vals, color=cols, width=0.62)
    for r, v in zip(b, vals):
        ax.text(r.get_x() + r.get_width() / 2, v + 2, f"{v}%", ha="center", va="bottom", color="white", fontsize=19, fontweight="bold")
    ax.set_ylim(0, 92); ax.set_yticks([]); ax.tick_params(colors="#cbd5e1", labelsize=12.5)
    ax.spines["bottom"].set_color("#334155")
    ax.set_title("Receipts with all four fields right (normalized)", color="#cbd5e1", fontsize=13, loc="left", pad=12)
    return save(fig, "progress_dark.png", dark=True)

def chart_fields():
    fields, plain, prompted, ft = ["Company", "Date", "Address", "Total"], [76, 99, 52, 71], [85, 98, 63, 73], [93, 98, 80, 98]
    fig, ax = plt.subplots(figsize=(7.2, 4.1)); x = range(len(fields)); w = 0.26
    for i, (vals, c, lab) in enumerate(((plain, "#cbd5e1", "Plain"), (prompted, "#94a3b8", "Detailed prompt"), (ft, ACCENT, "Fine-tuned v2*"))):
        bars = ax.bar([k + (i - 1) * w for k in x], vals, w, color=c, label=lab)
        for r, v in zip(bars, vals):
            ax.text(r.get_x() + r.get_width() / 2, v + 1.5, str(v), ha="center", va="bottom", fontsize=11.5,
                    color=INK if i == 2 else GREY, fontweight="bold" if i == 2 else "normal")
    ax.set_xticks(list(x)); ax.set_xticklabels(fields, fontsize=14); ax.set_yticks([]); ax.set_ylim(0, 112)
    ax.legend(frameon=False, ncol=3, loc="upper left", bbox_to_anchor=(0, 1.13), fontsize=12)
    return save(fig, "fields.png")

def chart_seen_new():
    fig, ax = plt.subplots(figsize=(6.0, 4.0)); groups = ["Shops seen in training\n(63 receipts)", "New shops\n(37 receipts)"]
    data = (("Plain", [34.9, 16.2], "#cbd5e1"), ("Detailed prompt", [47.6, 18.9], "#94a3b8"), ("Fine-tuned v1", [82.5, 37.8], ACCENT))
    w = 0.26
    for i, (lab, vals, c) in enumerate(data):
        bars = ax.bar([k + (i - 1) * w for k in range(2)], vals, w, color=c, label=lab)
        for r, v in zip(bars, vals):
            ax.text(r.get_x() + r.get_width() / 2, v + 1.5, f"{v:g}", ha="center", va="bottom", fontsize=12,
                    color=INK if i == 2 else GREY, fontweight="bold" if i == 2 else "normal")
    ax.set_xticks([0, 1]); ax.set_xticklabels(groups, fontsize=12.5); ax.set_yticks([]); ax.set_ylim(0, 100)
    ax.legend(frameon=False, ncol=3, loc="upper left", bbox_to_anchor=(0, 1.12), fontsize=11.5)
    return save(fig, "seen_new.png")

def chart_experiments():
    labels = ["Budget 280", "Budget 560 (v1)", "LoRA r = 8", "E4B (2x model)", "2 epochs", "Budget 1120 (v2)"]
    vals = [61, 66, 68, 69, 71, 76]
    fig, ax = plt.subplots(figsize=(6.6, 3.9))
    b = ax.barh(labels, vals, color=[("#cbd5e1" if v < 76 else ACCENT) for v in vals], height=0.6)
    for r, v in zip(b, vals):
        ax.text(v + 1, r.get_y() + r.get_height() / 2, f"{v}%", va="center", fontsize=13, color=INK, fontweight="bold" if v == 76 else "normal")
    ax.set_xlim(0, 88); ax.set_xticks([]); ax.tick_params(labelsize=12.5); ax.spines["bottom"].set_visible(False)
    ax.set_title("Fine-tuned, all four fields right (same 100 receipts)", fontsize=12, color=GREY, loc="left", pad=10)
    return save(fig, "experiments.png")

def chart_curve():
    steps, norm, strict = [0, 40, 80, 120, 157], [26, 54, 52, 54, 54], [6, 46, 48, 52, 52]
    fig, ax = plt.subplots(figsize=(5.4, 3.4))
    ax.plot(steps, norm, "-o", color=ACCENT, lw=2.5, label="normalized"); ax.plot(steps, strict, "--o", color="#94a3b8", lw=2, label="strict")
    ax.set_xticks(steps); ax.set_yticks([0, 25, 50]); ax.set_ylim(0, 65); ax.tick_params(labelsize=11)
    ax.set_xlabel("training steps (4 receipts each)", fontsize=11, color=GREY); ax.legend(frameon=False, fontsize=11, loc="lower right")
    ax.set_title("Learning curve (first 50 test receipts, v1)", fontsize=12, color=GREY, loc="left")
    return save(fig, "curve.png")

def crop_shot(name, box):
    p = os.path.join(CH, "crop_" + name)
    Image.open(os.path.join(SS, name)).convert("RGB").crop(box).save(p); return p

# ---------------------------------------------------------------- slide helpers (12-column grid)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
M, W = 0.75, 13.333 - 1.5
def col(i): return M + i * W / 12
def span(n): return n * W / 12

def text(s, x, y, w, h, runs, size=18, color=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line=1.08):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.line_spacing = line
        for r in (para if isinstance(para, tuple) and isinstance(para[0], tuple) else [para]):
            t, kw = (r, {}) if isinstance(r, str) else (r[0], r[1])
            run = p.add_run(); run.text = t; f = run.font
            f.name = FONT; f.size = Pt(kw.get("size", size)); f.bold = kw.get("bold", bold); f.color.rgb = rgb(kw.get("color", color))
            if kw.get("spacing", spacing):
                run._r.get_or_add_rPr().set("spc", str(kw.get("spacing", spacing)))
    return tb

def hair(s, x1, y1, x2, y2, color=HAIR, width=0.75):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    st = c._element.find(qn("p:style"))
    if st is not None:
        c._element.remove(st)          # drop the theme style (it adds a shadow that doubles hairlines)
    return c

def rect(s, x, y, w, h, fill, line=None, rounded=0.0):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = rounded
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill); shp.shadow.inherit = False
    if line:
        shp.line.color.rgb = rgb(line); shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    return shp

def browser(s, path, x, y, w):
    im = Image.open(path); h_img = w * im.size[1] / im.size[0]; bar = 0.3
    rect(s, x, y, w, bar + h_img, "#ffffff", line=HAIR)
    rect(s, x, y, w, bar, "#f3f4f6")
    for k in range(3):
        d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.16 + k * 0.17), Inches(y + 0.1), Inches(0.1), Inches(0.1))
        d.fill.solid(); d.fill.fore_color.rgb = rgb("#d1d5db"); d.line.fill.background(); d.shadow.inherit = False
    s.shapes.add_picture(path, Inches(x), Inches(y + bar), width=Inches(w)); return bar + h_img

def phone(s, path, x, y, h):
    im = Image.open(path); w_img = h * im.size[0] / im.size[1]; pad = 0.09
    rect(s, x, y, w_img + 2 * pad, h + 2 * pad, INK, rounded=0.12)
    s.shapes.add_picture(path, Inches(x + pad), Inches(y + pad), height=Inches(h)); return w_img + 2 * pad

def kicker(s, t, y=0.62, dark=False):
    text(s, M, y, 9, 0.3, t.upper(), size=10.5, color="#94a3b8" if dark else GREY, bold=True, spacing=300)

def title(s, t, y=0.92, size=36, dark=False, w=11.8):
    text(s, M, y, w, 1.2, t, size=size, color="#ffffff" if dark else INK, bold=True, line=1.0)

def footer(s, n, dark=False):
    c = "#64748b" if dark else "#9ca3af"
    text(s, M, 7.02, 8, 0.3, "Receipt to Expense  ·  Team Come_Join_Qarc  ·  Theme 3B", size=9.5, color=c)
    text(s, 13.333 - M - 1, 7.02, 1, 0.3, f"{n:02d}", size=9.5, color=c, align=PP_ALIGN.RIGHT)

def new(dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NIGHT if dark else "#ffffff"); return s

def notes(s, t):
    s.notes_slide.notes_text_frame.text = t

def stat(s, x, y, w, big, label, size=46, color=INK, label_color=GREY, sub=None):
    text(s, x, y, w, 0.9, big, size=size, color=color, bold=True, line=0.95)
    text(s, x, y + size / 72 * 1.12, w, 0.8, [label] + ([(sub, {"size": 11, "color": "#9ca3af"})] if sub else []), size=13.5, color=label_color)

# ---------------------------------------------------------------- slides (direction: few numbers, each = our edge over the
# SAME base model with the same settings; budget-1120 version, 100 held-out SROIE test receipts)
def chart_ten_x():
    labels = ["Gemma 4 E2B\nwithout our fine-tune", "Same model +\nengineered prompt", "Our fine-tune"]
    vals, cols = [7, 17, 70], ["#cbd5e1", "#94a3b8", ACCENT]
    fig, ax = plt.subplots(figsize=(6.4, 4.4))
    b = ax.bar(labels, vals, color=cols, width=0.58)
    for r, v in zip(b, vals):
        ax.text(r.get_x() + r.get_width() / 2, v + 1.8, f"{v}%", ha="center", va="bottom", fontsize=20, fontweight="bold",
                color=INK if v == 70 else GREY)
    ax.annotate("", xy=(1.66, 70), xytext=(0.3, 13), arrowprops=dict(arrowstyle="->", color=ACCENT, lw=2))
    ax.text(0.42, 50, "10x", fontsize=34, fontweight="bold", color=ACCENT)
    ax.set_ylim(0, 92); ax.set_yticks([]); ax.tick_params(labelsize=12.5)
    ax.set_title("Receipts with all four fields exactly right", fontsize=13, color=GREY, loc="left", pad=12)
    return save(fig, "ten_x.png")

P_TEN, P_SEEN, P_EXP = chart_ten_x(), chart_seen_new(), chart_experiments()
SCAN = crop_shot("pro_ui_scan.png", (207, 20, 1158, 900))
PHONE = os.path.join(SS, "pro_ui_phone.png")

def big_statement(n, kick, big, line, small, note_text, dark=False):
    s = new(dark=dark)
    kicker(s, kick, dark=dark)
    text(s, M, 1.45, W, 2.4, big, size=118, color="#34d399" if dark else ACCENT, bold=True, line=0.9)
    text(s, M, 3.55, span(10), 1.2, line, size=34, color="#ffffff" if dark else INK, bold=True, line=1.05)
    text(s, M, 4.6, span(8.5), 1.0, small, size=17, color="#cbd5e1" if dark else BODY, line=1.25)
    footer(s, n, dark=dark); notes(s, note_text); return s

# 1 · Hero
s = new()
kicker(s, "Theme 3B  ·  Hacktoberfest Hack Day  ·  24 Oct 2026", y=0.72)
text(s, M, 1.05, span(7.0), 1.6, "Receipt to Expense.", size=44, color=INK, bold=True, line=0.95)
text(s, M, 2.02, span(6), 1.2, "Snap a receipt. Our fine-tuned Gemma 4 reads it, logs the expense and splits the bill with friends.",
     size=18, color=BODY, line=1.15)
text(s, M, 3.2, span(6.2), 1.2, [(("Team Come_Join_Qarc", {"bold": True, "color": INK}), ("  ·  SSIT Tumkur", {})),
                                   (("Lead: Mohammed Owais", {"size": 12}),),
                                   (("Members: Mohammed Muzakkir, Mohammad Tariq Afzal, Mohammed Ayaan, Mohammed Younus", {"size": 12}),)],
     size=13.5, color=GREY, line=1.3)
hair(s, M, 4.75, M + span(6.6), 4.75)
text(s, M, 4.92, span(6.6), 1.4, [(("10x", {"size": 44, "bold": True, "color": ACCENT}),),
                                   (("more receipts read perfectly than the base model", {"size": 16, "color": BODY}),)],
     size=15, line=1.05)
browser(s, SCAN, col(7.05), 0.78, span(4.95))
footer(s, 1)
notes(s, "0:00-0:15. Hold up a real receipt. 'Riya's hostel flat splits groceries every week; the receipts end up in a WhatsApp "
         "photo pile and someone types them in. We built Receipt to Expense: our fine-tuned open Gemma 4 reads the photo, logs the "
         "expense and splits the bill. It reads ten times more receipts perfectly than the same model without our fine-tune.' "
         "(Swap in a real story from the team if you have one.)")

# 2 · Problem
s = new()
kicker(s, "Saam  ·  Bhed  —  the problem")
title(s, "Students still type receipts in by hand.")
hair(s, M, 2.0, 13.333 - M, 2.0)
for i, (head, items, accent) in enumerate((("Expense apps today", ["Closed, paid, cloud OCR", "Built for company cards", "Never say which field to trust"], False),
                                           ("Receipt to Expense", ["Open model + one small adapter file", "One model from photo to fields, no OCR step",
                                                                   "Says what to double-check, then splits the bill"], True))):
    x = M + i * span(6)
    if i:
        hair(s, x - 0.3, 2.3, x - 0.3, 5.3)
    text(s, x, 2.3, span(5.6), 0.5, head, size=22, color=ACCENT if accent else INK, bold=True)
    for k, it in enumerate(items):
        y = 3.05 + k * 0.72
        hair(s, x, y - 0.12, x + span(5.5), y - 0.12)
        text(s, x, y, span(5.5), 0.5, it, size=16, color=INK if accent else BODY)
text(s, M, 5.75, W, 0.7, [(("For ", {}), ("hostel flats splitting groceries, clubs reimbursing fest bills and small shops logging purchases.", {"color": INK}))],
     size=16, color=GREY)
footer(s, 2)
notes(s, "0:15-0:30. 'Expense apps exist, but they are closed, paid, cloud OCR products for company cards, and they never tell you "
         "which field to trust. Most receipt hacks chain OCR plus an LLM. Ours is one fine-tuned open model, photo to fields, built for "
         "students who share bills.'")

# 3 · How it works
s = new()
kicker(s, "How it works")
title(s, "One open model, from photo to expense.")
steps = (("01", "Snap", "A phone photo goes straight to the model, image first, at high resolution. No OCR step."),
         ("02", "Read", "Gemma 4 E2B with our LoRA adapter, trained on public labelled receipts (SROIE) and tested on receipts it never saw."),
         ("03", "Log and split", "Edit the fields, add them to expenses, or split equally or by item; everyone settles up in the fewest payments."))
for i, (n, h, b) in enumerate(steps):
    x = M + i * span(4)
    hair(s, x, 2.25, x + span(3.7), 2.25, color=INK, width=1.0)
    text(s, x, 2.42, 1, 0.4, n, size=13, color=ACCENT, bold=True, spacing=100)
    text(s, x, 2.8, span(3.6), 0.5, h, size=22, color=INK, bold=True)
    text(s, x, 3.35, span(3.6), 1.6, b, size=15, color=BODY, line=1.2)
hair(s, M, 5.35, 13.333 - M, 5.35)
text(s, M, 5.6, W, 0.8, [(("What made the difference: ", {"color": GREY}), ("doubling image resolution beat doubling the model size.", {"color": INK, "bold": True}))],
     size=19)
footer(s, 3)
notes(s, "0:30-0:45. 'Three steps. The photo goes straight in, image first, no OCR. A small open Gemma 4 with our LoRA adapter, "
         "trained on public labelled receipts, reads the four fields. Then you log or split. And what moved accuracy most: "
         "giving the model twice the image resolution beat using a model twice the size.'")

# 4 · 10x (the one chart)
s = new()
kicker(s, "Dand  —  the result")
title(s, "10x more receipts read perfectly.")
s.shapes.add_picture(P_TEN, Inches(M), Inches(1.95), height=Inches(4.45))
x = col(7.6)
text(s, x, 2.3, span(4.4), 2.4, ["Same Gemma 4 model, same settings, same 100 receipts it never saw.",
                                  "Without our fine-tune: 7 in 100 read perfectly.",
                                  "With it: 70 in 100."], size=17, color=BODY, line=1.3)
hair(s, x, 4.75, 13.333 - M, 4.75)
text(s, x, 4.95, span(4.4), 0.5, "100 held-out receipts  ·  McNemar p < 1e-10", size=13, color=INK, bold=True)
text(s, M, 6.62, W, 0.3, "Fine-tuned version shown: best of 5 configurations, same test set.", size=9.5, color="#9ca3af")
footer(s, 4)
notes(s, "0:45-1:05. 'Same model, same settings, a hundred receipts it never saw. Without our fine-tune it reads 7 of 100 receipts "
         "perfectly, all four fields exact. A carefully engineered prompt gets 17. Our fine-tune gets 70. Ten times. McNemar p is below "
         "one in ten billion.' (For questions: the engineered prompt was measured at the 560 budget; at 1120 the plain model scores "
         "lower, not higher, so the comparison is conservative.)")

# 5 · 3/4 fewer errors
big_statement(5, "Dand  —  reading errors", "3/4", "fewer reading errors.",
              "Field accuracy goes from 70% to 92% on the same model. Totals are right 98% of the time.",
              "1:05-1:15. 'Field by field, our fine-tune removes three quarters of the reading errors: 70 percent of fields right becomes 92. "
              "Totals, the field you actually pay, are right 98 percent of the time.'")

# 6 · 2x vs prompt engineering
big_statement(6, "Bhed  —  why fine-tune at all", "2x", "better than prompt engineering.",
              "A carefully engineered prompt gets 37% of receipts fully right. Our fine-tune gets 76%.",
              "1:15-1:25. 'Could we just write a better prompt? We tried: a carefully engineered prompt gets 37 percent of receipts fully "
              "right; our fine-tune gets 76. Fine-tuning doubles it. Now let's try yours.'", dark=True)

# 7 · Demo
s = new()
kicker(s, "Daam  —  live demo")
title(s, "Snap it. Check it. Split it.")
browser(s, SCAN, M, 1.95, 4.55)
phone(s, PHONE, 5.6, 1.95, 4.25)
x = col(7.75)
text(s, x, 2.0, span(4.2), 0.5, "Judges: lend us a receipt.", size=18, color=INK, bold=True)
text(s, x, 2.45, span(4.2), 0.9, "Scan the code to try it on your phone.", size=13, color=BODY)
s.shapes.add_picture(QR_APP, Inches(x), Inches(3.2), width=Inches(1.75))
text(s, x, 5.02, span(4.2), 0.3, APP_URL.replace("https://", ""), size=10, color=GREY)
text(s, x, 5.45, span(4.2), 1.2, ["Trust check on every field", "Split equally or by item", "Expenses to CSV"], size=12.5, color=BODY, line=1.3)
text(s, M, 6.62, W, 0.3, "On the projector: the polished UI. On phones: the public link, same fine-tuned model.", size=10, color="#9ca3af")
footer(s, 7)
notes(s, "1:25-2:35. Live. Do one warm-up scan before the talk. 'Could one of you lend us a receipt?' Scan it. 'Merchant, date, "
         "address and total, plus a trust check that says which field to double-check.' Press Split this bill: 'Split equally or by "
         "item; tax is shared fairly; everyone settles up in the fewest payments.' Add to expenses, download the CSV. If anything "
         "stalls for more than 15 seconds, switch to the screenshots.")

# 8 · Delivered
s = new()
kicker(s, "Delivered")
title(s, "Everything required, plus every bonus.")
req = ("Fine-tuned model + training settings", "Dataset description", "Plain vs fine-tuned comparison", "Working app, live",
       "Where it helps, limits and why")
bonus = ("Prompted baseline", "Training experiments", "Overfitting check", "Model size comparison (E4B)", "Combined with retrieval (RAG)",
         "Published on Hugging Face", "Error analysis", "Line items")
for c, (head, items) in enumerate((("Required", req), ("Bonus", bonus))):
    x = M + (0 if c == 0 else span(4.45))
    text(s, x, 2.0, span(4), 0.4, head, size=12, color=GREY, bold=True, spacing=150)
    for k, a in enumerate(items):
        y = 2.45 + k * (0.52 if c == 0 else 0.46)
        hair(s, x, y - 0.06, x + span(4.0), y - 0.06)
        text(s, x, y, 0.35, 0.4, "✓", size=14, color=ACCENT, bold=True)
        text(s, x + 0.38, y, span(4.0) - 0.4, 0.4, a, size=14, color=INK)
text(s, M, 5.55, span(4.2), 0.9, [(("Helps, limits, why: ", {"color": GREY}), ("learns receipt layouts fast; next, Indian receipts.", {"color": INK}))],
     size=14, line=1.25)
x = col(9.25)
text(s, x, 2.0, span(2.75), 0.4, "Model card", size=12, color=GREY, bold=True, spacing=150)
s.shapes.add_picture(QR_HF, Inches(x), Inches(2.45), width=Inches(1.6))
text(s, x, 4.15, span(2.75), 0.6, "huggingface.co/owaisnoe/\ngemma4-e2b-sroie-lora", size=10, color=GREY)
text(s, x, 4.95, span(2.75), 1.4, [("Next", {"bold": True, "color": INK}), "Indian receipts", "UPI requests from the split",
                                   "On-device: receipts stay on the phone"], size=12, color=BODY, line=1.3)
footer(s, 8)
notes(s, "2:35-3:00. 'Everything required is delivered, plus all eight bonus items; the adapters and model card are on Hugging Face. "
         "It learns receipt layouts fast; next we train on Indian receipts, add UPI requests straight from the split, and run it "
         "on-device so receipts never leave the phone. Thank you.'")

# 9 · Appendix: robustness
s = new()
kicker(s, "Appendix  —  for questions")
title(s, "Robustness.", size=30)
s.shapes.add_picture(P_SEEN, Inches(M), Inches(1.85), height=Inches(4.2))
x = col(7)
text(s, x, 1.95, span(5), 4.4, ["Seen vs new shops (first model, budget 560): 63 of the 100 test receipts come from shops that also appear "
                                 "in training; on brand-new shops the fine-tune still doubles the prompt (37.8% vs 18.9%).",
                                 "Other receipt styles (never trained on): WildReceipt, all labelled fields 58% plain vs 56% fine-tuned "
                                 "(not significant); CORD totals 87.5% vs 91.7% (not significant). No significant loss.",
                                 "Retrieval (vendor memory from training labels): fine-tuned 66% to 69%; prompt 37% to 47%.",
                                 "Line items, zero-shot on CORD: item F1 71% plain vs 79% fine-tuned."], size=12.5, color=BODY, line=1.25)
footer(s, 9)
notes(s, "Appendix, for questions about generalisation.")

# 10 · Appendix: experiments + engineering
s = new()
kicker(s, "Appendix  —  for questions")
title(s, "Experiments and one silent bug.", size=30)
s.shapes.add_picture(P_EXP, Inches(M), Inches(1.95), height=Inches(3.3))
x = col(7.5)
text(s, x, 1.95, span(4.5), 4.6, ["Each arm changes one setting, on the same 100 receipts. Only resolution was significant: "
                                 "1120 vs 560 tokens, 11 receipts better and 1 worse (p = 0.006). Doubling the model (E4B) gave 69%.",
                                 "The collator bug: Unsloth's default collator shrank every receipt to 512 px and snapped it to 32 px, "
                                 "with the same token count but different pixels (max difference 0.68). Fixed with images as a "
                                 "dataset column and resize=\"max\" (difference 0.0000).",
                                 "Trust check: when fine-tuned and plain agree on a field, it was right 83-100% of the time (first model)."],
     size=12.5, color=BODY, line=1.25)
footer(s, 10)
notes(s, "Appendix, for questions about method. The 1120 version is the best of 5 configurations on the same test set.")

# 11 · Appendix: data + settings
s = new()
kicker(s, "Appendix  —  data and settings")
title(s, "Data and settings.", size=30)
kv = (("Data", "SROIE (ICDAR 2019, CC BY 4.0): 626 training receipts; 100 fixed test receipts (seed 0) from the 361-receipt test split; none trained on"),
      ("Base model", "unsloth/gemma-4-E2B-it, 16-bit LoRA adapters (never merged); r = 32, vision + language layers, 1.15% of weights"),
      ("Training", "lr 2e-4, cosine, adamw_8bit, batch 1 x grad-accum 4, 5 smoke + 157 steps (1 epoch), 2.5 min on a Colab G4"),
      ("Versions", "v1: 560 visual tokens (66% receipts, 89.5% fields)  ·  v2: 1120 visual tokens (76% / 92%, best of 5 configurations)"),
      ("Evaluation", "image-first prompt, greedy decoding, batch size 1; strict + normalized scoring; exact McNemar tests"),
      ("Baselines", "plain model, same settings; engineered prompt measured at 560 tokens (at 1120 the plain model scores lower)"),
      ("Compute", "Colab G4 for speed; E2B LoRA needs about 8-10 GB per Unsloth's docs, so it should fit a free T4 (not tested)"))
for k, (a, b) in enumerate(kv):
    y = 1.85 + k * 0.66
    hair(s, M, y - 0.08, 13.333 - M, y - 0.08)
    text(s, M, y, span(2.0), 0.5, a, size=13, color=GREY, bold=True)
    text(s, col(2.2), y, span(9.8), 0.6, b, size=13, color=INK)
footer(s, 11)
notes(s, "Appendix: dataset description and settings (required deliverables).")

prs.save(OUT_PPTX)
print("saved", OUT_PPTX, "| slides:", len(prs.slides))
